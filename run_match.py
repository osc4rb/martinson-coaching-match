"""One-command matching: Forms export .xlsx -> Matches.xlsx in the OneDrive folder the flow reads.

Usage:
  python run_match.py <responses.xlsx> [--coach 2028,2029] [--coachee 2030] [--out PATH]
Default roles: youngest class is coached, every older class coaches.
"""
import argparse
import json
import subprocess
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = Path(__file__).parent
ONEDRIVE = Path.home() / "OneDrive - USAFA (AFAcademy.AF.edu)" / "Martinson Coaching Match" / "Matches.xlsx"
COLS = ["Coach", "Coach Class", "Coach Email", "Coachee", "Coachee Class", "Coachee Email",
        "Score", "Shared Interests", "Intro Message", "Teams Link"]


def read_first_sheet(path):
    ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = [["" if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(c.strip() for c in r)]
    return {"headers": rows[0], "rows": rows[1:]}


def write_matches(pairs, out):
    wb = Workbook()
    ws = wb.active
    ws.title = "Matches"
    ws.append(COLS)
    for p in pairs:
        ws.append([p["coach"], p["coachClass"], p["coachEmail"], p["coachee"], p["coacheeClass"],
                   p["coacheeEmail"], p["score"], p["shared"], p["message"], p["link"]])
    t = Table(displayName="MatchTable", ref=f"A1:J{len(pairs) + 1}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    for col, w in zip("ABCDEFGHIJ", [24, 11, 34, 24, 13, 34, 8, 70, 70, 40]):
        ws.column_dimensions[col].width = w
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("responses")
    ap.add_argument("--coach")
    ap.add_argument("--coachee")
    ap.add_argument("--out", default=str(ONEDRIVE))
    a = ap.parse_args()

    tmp = Path(tempfile.mkdtemp())
    (tmp / "in.json").write_text(json.dumps(read_first_sheet(a.responses)), encoding="utf-8")
    cmd = ["node", str(HERE / "match_cli.cjs"), str(tmp / "in.json"), str(tmp / "out.json")]
    if a.coach or a.coachee:
        cmd += [a.coach or "", a.coachee or ""]
    subprocess.run(cmd, check=True)
    result = json.loads((tmp / "out.json").read_text(encoding="utf-8"))
    if not result["pairs"]:
        raise SystemExit("No pairs made: check class roles.")
    write_matches(result["pairs"], Path(a.out))
    print(f"wrote {a.out}")
    for p in result["pairs"]:
        print(f"  {p['coach']} ({p['coachClass']}) -> {p['coachee']} ({p['coacheeClass']})  score {p['score']}  | {p['shared'][:70]}")
    if result["idle"]:
        print("  coaches without a coachee:", ", ".join(result["idle"]))


if __name__ == "__main__":
    main()
