"""Generates the form question sheet, the blank template workbook, and a fake
sample workbook for testing. Run: python build_files.py"""
import json
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = Path(__file__).parent

# (column header, question text, type, options, scored?)
QUESTIONS = [
    ("Name", "Full name", "Text", None, False),
    ("Email", "USAFA email", "Text", None, False),
    ("Class Year", "Class year", "Choice", ["2027", "2028", "2029", "2030"], False),
    ("Squadron", "Squadron (number)", "Text", None, False),
    ("Coaching Focus",
     "What do you want help with, or want your coach to focus on? (pick up to 3)",
     "Checkbox (pick up to 3)",
     ["Academics & study habits", "Choosing a major", "Fellowships & scholarships",
      "Grad school planning", "AF/SF career planning", "Research opportunities",
      "Navigating cadet life", "Leadership & military training", "Time management",
      "Wellness & balance"], True),
    ("Major", "Major (if undeclared: majors you're considering)", "Checkbox (pick up to 3)",
     ["Aeronautical Engineering", "Astronautical Engineering", "Behavioral Sciences",
      "Biology", "Chemistry", "Civil Engineering", "Computer Science", "Cyber Science",
      "Data Science", "Economics", "Electrical & Computer Engineering", "English",
      "Foreign Area Studies", "Geospatial Science", "History", "Humanities",
      "Legal Studies", "Management", "Mathematics", "Mechanical Engineering",
      "Meteorology", "Military & Strategic Studies", "Operations Research",
      "Philosophy", "Physics", "Political Science", "Space Operations",
      "Systems Engineering", "Undecided"], True),
    ("Minor", "Minor(s) you have or are considering", "Checkbox",
     ["Foreign language", "Philosophy", "Religion Studies", "Space Studies",
      "Nuclear Weapons & Strategy", "Global Logistics", "Other", "None / undecided"], True),
    ("Grad School Plans", "Grad school / fellowships you're interested in", "Checkbox",
     ["Rhodes", "Marshall", "Fulbright", "Truman", "Schwarzman", "Gates Cambridge",
      "Mitchell", "Hertz / NSF GRFP", "Master's right after graduation",
      "PhD", "Medical school", "Law school", "Not sure yet", "Not planning grad school"], True),
    ("AF Career Interests", "Air Force / Space Force career fields you're interested in",
     "Checkbox (pick up to 4)",
     ["Pilot", "Combat Systems Officer", "RPA Pilot", "Air Battle Manager",
      "Space Operations", "Cyber", "Intelligence", "Developmental Engineer",
      "Scientist", "Acquisitions / Program Management", "Special Warfare", "Missiles",
      "Maintenance / Logistics", "Security Forces", "OSI", "Medical", "JAG",
      "Weather", "Finance / Contracting", "Public Affairs", "Undecided"], True),
    ("Future Path", "Long-term (after your commitment), where do you see yourself?",
     "Checkbox",
     ["Career military", "Industry / tech", "Academia / research",
      "Government / policy", "Entrepreneurship", "Medicine", "Law", "Not sure yet"], True),
    ("Opportunities", "Opportunities you want to pursue at USAFA", "Checkbox",
     ["Summer research", "Internships", "Study abroad / exchange semester",
      "Ops Air Force", "Soaring / powered flight", "Jump", "Cadet leadership positions",
      "Competition teams / clubs", "Publishing research", "Conferences / presenting"], True),
    ("Research Areas", "Research areas you're interested in", "Checkbox (pick up to 4)",
     ["AI / machine learning", "Autonomy / drones", "Aerospace / hypersonics", "Space",
      "Cybersecurity", "Energy", "Biomedical / human performance", "Climate / environment",
      "Quantum / physics", "Materials", "National security / policy", "History",
      "Ethics / philosophy", "Economics", "Psychology / leadership", "Education"], True),
    ("Research Interests", "In a sentence or two, what topics are you curious about?",
     "Long text", None, True),
    ("Fun", "What do you do for fun?", "Checkbox",
     ["Running", "Lifting", "Climbing", "Hiking / outdoors", "Skiing / snowboarding",
      "Team sports", "Music", "Gaming", "Reading", "Writing", "Art / photography",
      "Cooking", "Faith community", "Volunteering", "Travel", "Film / TV", "Chess / board games"], True),
    ("Anything Else", "Anything else your coach/coachee should know? (not used for matching)",
     "Long text", None, False),
]
HEADERS = [q[0] for q in QUESTIONS]


def write_form_sheet():
    lines = [
        "# Martinson Coaching Match: Microsoft Forms questions",
        "",
        "Build ONE form for both classes. Type options **exactly** as written. The matcher",
        "compares option text, so a typo in one place makes two identical answers look different.",
        "",
        "Settings (... > Settings): *Only people in my organization can respond*, *Record name*,",
        "*One response per person*.",
        "",
    ]
    for i, (header, text, kind, options, scored) in enumerate(QUESTIONS, 1):
        tag = "" if scored else "  _(not scored)_"
        lines.append(f"### {i}. {text}{tag}")
        lines.append(f"- Type: **{kind}**. Required: {'yes' if header != 'Anything Else' else 'no'}")
        lines.append(f"- Excel column: `{header}`")
        if options:
            lines.append("- Options:")
            lines.extend(f"  - {o}" for o in options)
        lines.append("")
    (OUT / "form-questions.md").write_text("\n".join(lines), encoding="utf-8")


def make_workbook(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Responses"
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    if not rows:
        ws.append([""] * len(HEADERS))  # Excel tables need at least one body row
    ref = f"A1:{ws.cell(row=ws.max_row, column=len(HEADERS)).coordinate}"
    t = Table(displayName="Responses", ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    for c in ws[1]:
        c.font = Font(bold=True)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 22
    wb.save(path)


def sample_rows(n28=4, n29=10, n30=14, seed=7):
    rnd = random.Random(seed)
    opts = {q[0]: q[3] for q in QUESTIONS if q[3]}
    first = ["Alex", "Jordan", "Sam", "Riley", "Taylor", "Morgan", "Casey", "Jamie", "Avery",
             "Quinn", "Drew", "Reese", "Skyler", "Parker", "Rowan", "Emerson", "Hayden",
             "Kendall", "Logan", "Blake", "Cameron", "Dakota", "Elliot", "Finley"]
    last = ["Nguyen", "Smith", "Garcia", "Patel", "Kim", "Johnson", "Lopez", "Brown",
            "Davis", "Martinez", "Wilson", "Anderson", "Thomas", "Moore"]
    topics = ["swarm autonomy for drones", "hypersonic flow and heat shields",
              "machine learning for satellite imagery", "ethics of autonomous weapons",
              "Cold War nuclear strategy", "human performance and sleep",
              "quantum sensing", "cyber defense of space systems", "great power competition",
              "leadership psychology under stress", "renewable energy on bases"]
    rows = []
    for i in range(n28 + n29 + n30):
        cls = "2028" if i < n28 else "2029" if i < n28 + n29 else "2030"
        name = f"{first[i % len(first)]} {last[(i * 5 + i // len(first)) % len(last)]}"
        email = f"C{cls[2:]}{name.replace(' ', '.')}@example.com"
        row = []
        for h in HEADERS:
            if h == "Name": row.append(name)
            elif h == "Email": row.append(email)
            elif h == "Class Year": row.append(cls)
            elif h == "Squadron": row.append(str(rnd.randint(1, 40)))
            elif h == "Research Interests": row.append(f"I'm interested in {rnd.choice(topics)} and {rnd.choice(topics)}.")
            elif h == "Anything Else": row.append("")
            else:
                picks = rnd.sample(opts[h], k=min(len(opts[h]), rnd.randint(1, 3)))
                # Mix both formats the script must handle
                row.append(json.dumps(picks) if i % 2 else ";".join(picks) + ";")
        rows.append(row)
    return rows


if __name__ == "__main__":
    write_form_sheet()
    make_workbook(OUT / "Martinson_Matching_TEMPLATE.xlsx", [])
    rows = sample_rows()
    make_workbook(OUT / "Martinson_Matching_SAMPLE.xlsx", rows)
    (OUT / "test" ).mkdir(exist_ok=True)
    (OUT / "test" / "sample.json").write_text(json.dumps({"headers": HEADERS, "rows": rows}, indent=1))
    print("wrote form-questions.md, TEMPLATE.xlsx, SAMPLE.xlsx, test/sample.json")
